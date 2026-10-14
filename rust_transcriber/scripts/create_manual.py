#!/usr/bin/env python3
"""Generate the Word manual; requires python-docx and docs/assets screenshots."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
d=Document(); sec=d.sections[0]
sec.page_height=Inches(11.7);sec.page_width=Inches(8.3)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.75)
normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',34),('Heading 1',22),('Heading 2',14)]:
 s=d.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('2E5266')
header=sec.header.paragraphs[0];header.text='DAS TRANSCRIBER  /  USER MANUAL';header.style='Caption'
f=sec.footer.paragraphs[0];f.text='Version 2.0.0 • 14 October 2026                                      Page '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
d.core_properties.title='DAS Transcriber — User Manual'
d.core_properties.subject='HDF5 to SEG-Y / SEG-D: command line and PyQt6 GUI'
d.core_properties.author='DAS project'
def p(t='',style=None):return d.add_paragraph(t,style)
def h(t):d.add_heading(t,2)
def page(t):d.add_page_break();d.add_heading(t,1)
def code(t):
 x=p();x.paragraph_format.space_before=Pt(4);x.paragraph_format.space_after=Pt(8)
 for i,line in enumerate(t.splitlines()):
  r=x.add_run(('\n' if i else '')+line);r.font.name='Consolas';r.font.size=Pt(8)
 shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEF2F5');x._p.get_or_add_pPr().append(shade)
def table(headers,rows):
 t=d.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for row in t.rows:
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  for c in row.cells:
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(4)
    for r in par.runs:r.font.size=Pt(9)
def step(t):p(t,'List Number')
d.add_heading('DAS Transcriber',0)
p('User manual', 'Subtitle')
p('HDF5 → SEG-Y revision 1 / SEG-D revision 3.1', 'Subtitle')
p('Command-line converter and PyQt6 desktop GUI\nVersion 2.0.0 • 14 October 2026')
h('What this tool does')
p('Converts the project’s DAS HDF5 recordings into seismic files with receiver coordinates. Choose continuous records grouped into one file per UTC day, or shot windows defined by supported P1/11 navigation files. All selected channels are included in each applicable output file.')
p('The seismic writers and navigation parser are implemented in this project. RustSU and external SEG-D libraries are not required. Conversion uses a receiver CSV and does not connect to the database.')
h('Choose the correct duration control')
table(['Mode','GUI control','CLI option','Default'],[
('Continuous','Continuous trace length (seconds)','--trace-length','60 s'),
('Shot','Shot duration (seconds)','--duration','GUI: 10 s; CLI: required'),
('Shot timing','Shot start offset (seconds)','--offset','0 s')])
p('Example: a 10-second shot with offset −1 starts one second before the shot and ends approximately nine seconds after it. Durations and starts align to acquired samples.')
h('Guide contents')
for t in ['1. Installation and first run','2. Inputs and receiver positions','3. GUI: continuous conversion','4. GUI: shot slicing','5. Command-line recipes','6. Output files and quality checks','7. Troubleshooting','8. CLI argument reference','9. Input and format reference']:p(t)
p('Examples use illustrative paths. Replace them with paths accessible on the computer running the converter. Screenshots show example settings, not validated conversion results. Use Word’s Navigation Pane to jump between headings.')
page('1. Installation and first run')
p('The current implementation targets Linux. Run the following commands from the upgraded-guacamole project directory. The GUI starts the Rust executable on the same computer; it does not submit remote SSH jobs.')
h('Requirements')
table(['Component','Requirement'],[('Rust','Version 1.88 or newer, with Cargo'),('Native build tools','C compiler, pkg-config, HDF5 development headers'),('Debian/Ubuntu packages','build-essential pkg-config libhdf5-dev'),('Desktop GUI','Python 3, venv support, PyQt6 and a graphical desktop'),('UTM navigation only','PROJ cs2cs command (Debian/Ubuntu: proj-bin)')])
h('Build the converter')
code('cargo build --release --locked \\\n  --manifest-path rust_transcriber/Cargo.toml')
h('Prepare and launch the GUI')
code('python3 -m venv rust_transcriber/.venv\nrust_transcriber/.venv/bin/python -m pip install \\\n  -r rust_transcriber/requirements-gui.txt\nrust_transcriber/.venv/bin/python rust_transcriber/gui.py')
p('If the environment is already prepared, run only the final launch command. The GUI selects target/release/rust_transcriber when present, otherwise target/debug/rust_transcriber. Confirm the CLI executable field points to the intended build.')
h('Verify the CLI')
code('rust_transcriber/target/release/rust_transcriber --version\nrust_transcriber/target/release/rust_transcriber --help')
h('Working with server recordings')
p('Mount server storage locally for GUI access, copy the desired files locally, or run the CLI on the server. A server path is usable only on a machine where that path exists. Do not enter an SSH URL in the H5 inputs box.')
p('The server archive follows this directory structure (the spelling is intentional):')
code('/raid1/fsi/exps/\n  Leknse_Solbjorn_Kvitfoss_Mux_Phase/YYYYMMDD/dphi/')
p('The deployed project is /home/inst/oistein/upgraded-guacamole on server 10.147.20.10. From that project directory, the relative CLI commands in this guide apply. Use completed, stable recordings and a separate output directory with sufficient free space.')
page('2. Inputs and receiver positions')
h('Select HDF5 recordings')
p('Supply individual .h5/.hdf5 files, directories scanned recursively, or glob patterns. In the GUI, use one path or pattern per line. In the CLI, repeat --input and quote glob patterns. Files are deduplicated by canonical path and sorted by acquisition time. Directory symlinks encountered while scanning are skipped.')
h('Select native channels')
code('--channels "1-130,50000-78000"')
p('Ranges include both endpoints. Values are native IDs from header/channels, not array columns, CSV row numbers, mast IDs, or group IDs. The default all selects every recorded channel. Overlapping ranges are merged; absent IDs inside a range are not created.')
p('Use the file’s explicit channel IDs for both older 4 m and newer 8 m acquisition configurations. Do not multiply channel IDs to compensate for spatial decimation. Inspect the selected channel count and IDs during Check inputs.')
h('Export receiver positions from the server')
code('bash rust_transcriber/scripts/export_receivers.sh > receivers.csv')
p('This script uses SSH to inst@10.147.20.10 and a read-only database transaction to export public.masteliste. Its optional first argument overrides the SSH destination. Conversion itself needs only the resulting CSV.')
table(['CSV column','Meaning'],[('lid','DAS line ID; must match the selected --lid'),('channel','Exact native HDF5 channel ID'),('latitude','WGS84 latitude in decimal degrees'),('longitude','WGS84 longitude in decimal degrees'),('distance, gid and other columns','Accepted additional metadata; not interpreted for conversion')])
p('Every selected recorded channel requires an exact (lid, channel) match with valid coordinates. Duplicate keys are invalid. A full server export can still miss channels when acquisition and geometry sampling phases differ. Provide a verified matching CSV or narrow the selection; the tool does not interpolate coordinates. A mast-only CSV is unsuitable.')
h('Restrict acquisition time')
code('--start 2026-09-24T00:00:00Z \\\n--end   2026-09-25T00:00:00Z')
p('Start is inclusive; end is exclusive. Use RFC3339 with an explicit timezone, such as Z for UTC or +02:00. Daily output is always grouped by UTC date, even when selection times use another offset. Leave both fields blank to use the available input coverage.')
page('3. GUI: continuous conversion')
d.add_picture(str(ROOT/'docs/assets/gui-continuous.png'),width=Inches(5.7))
p('Figure 1. Continuous mode with illustrative input and output paths.', 'Caption')
step('Select H5 inputs, Receiver CSV, Native channels and DAS line ID. Choose Mode = continuous and Output format = segy, segd or both.')
step('Set Continuous trace length (seconds). The default is 60. Set optional Start and End bounds, then choose a new Output directory.')
step('Click Check inputs. Review selected channels, planned filenames, estimated output size and warnings. A successful check enables Start conversion.')
step('Click Start conversion and monitor the log. Keep the adjacent JSON reports with the output files.')
p('One file per UTC day contains all selected channels and many timed trace blocks. The trace length is not the daily file length. Blocks can be shorter at gaps, midnight, acquisition changes or selection boundaries. Compatible consecutive H5 files are stitched without filling gaps.')
page('4. GUI: shot slicing')
d.add_picture(str(ROOT/'docs/assets/gui-shot.png'),width=Inches(5.7))
p('Figure 2. A 10-second shot window beginning one second before the shot.', 'Caption')
step('Select the H5 inputs and receiver settings, then choose Mode = shot. Add the P1/11 navigation files; each basename must be unique.')
step('Set P1/11 source XY coordinates: utm33n for EPSG:32633 easting/northing, or wgs84 for longitude/latitude in degrees.')
step('Set Shot duration (seconds) and Shot start offset (seconds). For a 10-second window starting one second early, use 10 and −1. Continuous trace length is disabled in this mode.')
step('Choose format and output directory. Click Check inputs, review skipped shots and warnings, then Start conversion.')
p('The first acquired sample at or after shot time + offset is selected. Duration rounds to the nearest whole sample. Incomplete windows and windows crossing acquisition configuration changes are skipped and reported. Each navigation input produces one file per selected format containing its complete shots.')
page('5. Command-line recipes')
p('Run these commands from the project directory. Replace /data paths with your own. Begin with --dry-run to validate without writing outputs; remove it to convert. The GUI’s Equivalent command provides the command for the current settings.')
h('Continuous: 60-second traces, one file per UTC day')
code('rust_transcriber/target/release/rust_transcriber \\\n  --input /data/das/20260924/dphi \\\n  --receivers-csv receivers.csv --lid 1 \\\n  --channels "1-130,50000-78000" \\\n  --mode continuous --trace-length 60 \\\n  --format both --output /data/exports/day24 --dry-run')
p('Change --trace-length to select another continuous trace duration. Repeat --input to include additional files or directories. A multi-day selection produces separate files for each UTC day with usable data.')
h('Shot: 10 seconds, starting 1 second before the shot')
code('rust_transcriber/target/release/rust_transcriber \\\n  --input /data/das/20260924/dphi \\\n  --receivers-csv receivers.csv --lid 1 --channels "0-128" \\\n  --mode shot --p111 "/data/navigation/*.p111" \\\n  --p111-crs utm33n --duration 10 --offset -1 \\\n  --format both --output /data/exports/shots --dry-run')
p('For 20 seconds starting exactly at the shot, use --duration 20 --offset 0. For WGS84 source coordinates use --p111-crs wgs84. The CLI requires --duration in shot mode; it does not inherit the GUI’s 10-second default.')
h('Saving and reusing GUI jobs')
p('Save job… stores the GUI settings as JSON; Load job… restores them. This stores configuration, not recordings or results. Check executable, input, receiver and output paths after moving a job to another machine. Run Check inputs again after loading or editing settings.')
h('Cancelling a conversion')
p('Click Cancel in the GUI, or interrupt the CLI normally. The GUI requests graceful termination and waits for cleanup. Closing the GUI during a job initiates cancellation. Completed outputs from earlier in the job remain; active temporary files are removed on normal cancellation.')
p('A forced kill or power loss can leave .partial files. Inspect these before removing or relocating them. The converter refuses existing output names and partial files; use a fresh output directory for a new attempt.')
page('6. Output files and quality checks')
table(['Mode','Example output','Contents'],[('Continuous','20260924_lid1.sgy\n20260924_lid1.segd','All selected channels for one UTC day, divided into timed blocks'),('Shot','survey.sgy\nsurvey.segd','Complete shots from survey.p111'),('Report','20260924_lid1.sgy.json\nsurvey.segd.json','Settings, geometry, provenance, timing, gaps and skipped shots')])
h('What happens to the samples')
p('Output samples are big-endian IEEE float32 values calculated as H5 data × header/dataScale. The physical unit is retained in JSON. The converter performs no filtering, resampling, rotation, strain conversion or clock correction. Nonfinite values and float32 overflow fail conversion.')
p('Requested durations round to whole samples. For an interval of 0.00204 seconds, a 60-second request gives 29,412 samples, or 60.000480 seconds. Reports retain actual timing and the mapping back to source files and sample ranges.')
h('Choose a compatible trace length')
p('SEG-Y revision 1 supports at most 65,535 samples per trace and a sample interval of at most 65,535 microseconds. Requested duration is checked even when available data would create shorter traces. Selecting both formats still enforces SEG-Y limits.')
table(['Sample interval','Requested duration','Result'],[('2.04 ms','60 seconds','29,412 samples: within the sample-count limit'),('0.5 ms','60 seconds','120,000 samples: too long for this SEG-Y writer'),('0.5 ms','32 seconds','64,000 samples: within the sample-count limit')])
p('Shorten --trace-length (continuous) or --duration (shot), or select SEG-D. The supported SEG-D normal-record duration is approximately 2,147 seconds maximum; preflight checks the actual request.')
h('Review every finished job')
step('Confirm successful completion in the GUI log or CLI exit code. Check all planned outputs and their adjacent JSON reports exist.')
step('Review warnings, excluded files, gaps and skipped shots. A successful conversion does not imply uninterrupted coverage or verified clock synchronization.')
step('Check native channel IDs, receiver positions, units and actual timing against the source data. For shots, compare original shot time and actual trace start.')
step('Open a representative file in the downstream seismic application before a large production run. SEG-Y files can contain variable-length traces; use a reader that supports them.')
p('Files and reports are individually published without overwriting existing files; the whole job is not one atomic transaction. Independent decoding tests cover SEG-D, but downstream SEG-D application interoperability and production navigation-file validation remain necessary.')
page('7. Troubleshooting')
table(['Symptom','Action'],[
('Start conversion is disabled','Run Check inputs successfully. Any setting change invalidates the previous check.'),
('CLI will not start','Check the executable path and execute permission. Build the release executable; inspect the log for missing native libraries.'),
('No input files / inaccessible path','Verify paths on the machine running the converter. Mount server storage first. Quote CLI globs; put one GUI path per line.'),
('Missing receiver coordinates','Check the line ID and exact native channel ID in the CSV. Use a verified matching geometry file or narrow the selection.'),
('SEG-Y sample-count limit','Reduce the active duration control or choose segd. Selecting both still requires SEG-Y-compatible lengths.'),
('Input file excluded for quality metadata','Inspect missingSamples and source quality. Missing or unknown quality metadata excludes the entire affected file; gaps are not padded.'),
('Overlapping recordings rejected','Remove conflicting copies or overlapping selections. Canonical-path deduplication does not resolve different files with conflicting coverage.'),
('Shots skipped / no complete shots','Check H5 coverage, offsets, time bounds, gaps and configuration changes. A shot must have a complete compatible window.'),
('Unsynchronized clock','Explicit timeSynchronized=0 prevents shot extraction. Verify acquisition clock status and the instrument’s status convention.'),
('Navigation parse or CRS error','Confirm the supported legacy S-record layout and timestamp format. Choose the correct CRS; install proj-bin for utm33n.'),
('Output already exists / .partial exists','Use a new output directory. Inspect previous outputs and interrupted partials before manually moving or removing them.'),
('Downstream SEG-Y reader fails','Boundary traces may have different lengths. Use a reader supporting variable-length revision 1 files; some readers assume fixed lengths.'),
('Input changed during conversion','Use completed, stable H5 files. Source size and modification time are checked before and after reads.')])
h('Logs and exit codes')
p('CLI stdout contains one JSON event per line, including indexing, warning, preflight, progress, output_complete, complete and error. Preserve the log and adjacent reports when diagnosing a job.')
table(['Exit code','Meaning'],[('0','Success'),('1','Validation or conversion error'),('2','Invalid command-line arguments'),('130','Cancelled')])
page('8. CLI argument reference')
p('Required common arguments: --input, --receivers-csv, --lid and --output. Shot mode additionally requires --p111, --p111-crs and --duration.')
table(['Argument','Use / default'],[
('--input PATH','Repeatable file, recursive directory or quoted glob; required.'),
('--receivers-csv PATH','Receiver CSV with lid, channel, latitude, longitude; required.'),
('--lid INTEGER','DAS line ID matching the receiver CSV; required. Range 1–8,388,607.'),
('--channels SELECTION','all (default), or comma-separated native IDs and inclusive ranges.'),
('--mode MODE','continuous (default) or shot.'),
('--format FORMAT','segy (default), segd or both.'),
('--output DIRECTORY','Destination directory; required. Existing output names are refused.'),
('--trace-length SECONDS','Continuous mode only; default 60; rounded to a whole sample.'),
('--p111 PATH','Shot mode navigation; repeatable path or glob.'),
('--p111-crs CRS','Required in shot mode: utm33n or wgs84.'),
('--duration SECONDS','Required shot duration; no CLI default. GUI starts at 10 seconds.'),
('--offset SECONDS','Shot start relative to shot time; default 0; negative allowed.'),
('--start TIMESTAMP','Optional inclusive RFC3339 start with an explicit timezone.'),
('--end TIMESTAMP','Optional exclusive RFC3339 end with an explicit timezone.'),
('--dry-run','Validate and print the plan without creating outputs.'),
('--help / --version','Display usage / executable version.')])
h('Timing rules at a glance')
p('Continuous mode uses --trace-length. Shot mode uses --duration and --offset. Do not use --trace-length to set shot duration. The GUI disables the inactive timing controls and passes only options relevant to the selected mode.')
p('Continuous gaps remain missing timed blocks. Complete shot windows must fit compatible acquired data; incomplete windows are skipped rather than padded or shortened. If no complete output blocks remain, validation fails.')
page('9. Input and format reference')
h('Supported HDF5 profile')
p('The reader expects numeric data as samples × channels, a unique integer header/channels vector, Unix-seconds header/time, header/dt in seconds representable as integer microseconds, and scaling/unit metadata. missingSamples must be present and empty for usable files. Configuration boundaries include the full native channel grid, ROI metadata, gauge length, scaling, units and selected columns.')
h('Supported legacy P1/11 profile')
p('The parser accepts the original program’s comma-separated S-record layout with a header record. It does not implement every P1/11 variant or general CRS record routing. Validate a representative navigation file first.')
table(['One-based column','Value'],[('4','Integer source line'),('6 / 7','Integer shot point / source point index'),('8','UTC shot timestamp'),('10','Source identifier: final decimal digit'),('13 / 14','Source X / Y'),('15 / 27','Source depth / water depth, metres')])
p('Accepted timestamps: YYYY:MM:DD:HH:MM:SS[.ffffff] or YYYY:DDD:HH:MM:SS[.ffffff]. Fractions are decimal fractions of a second. Leap-second source records are rejected. wgs84 means X=longitude, Y=latitude; utm33n means EPSG:32633 easting/northing, transformed through PROJ.')
h('SEG-Y field conventions')
table(['One-based trace-header bytes','Encoding'],[('73–88','Source/receiver X,Y: arcseconds, coordinate scalar −1000'),('157–168','Actual trace-start UTC calendar fields; time basis 4'),('189–192 / 193–196','DAS line ID / source line ID (zero for continuous)'),('229–232','Custom native channel ID, unsigned big-endian integer'),('233–240','Custom actual start, signed Unix microseconds')])
p('The file has a 3,200-byte ASCII text header, 400-byte binary header and 240-byte trace headers. The fixed-length flag is zero; each trace carries its own count and interval. Continuous source coordinates and shot point are zero/undefined. Exact shot identity, time and offsets remain in JSON.')
h('SEG-D conventions')
p('Revision 3.1 output uses a 128-byte storage-unit label, multiple records, sample format 8058, descale factor 1 and receiver CRS EPSG:4326. Physical units are preserved in channel-set text and JSON; the numeric unit code is unknown where no matching DAS unit exists. Timestamps use GPS microseconds from 1980-01-06 with a leap-second table through 2017-01-01. Earlier dates and leap-insertion-crossing records are unsupported; future announced leap seconds require a table update.')
p('Implementation reference: rust_transcriber/README.md, src/args.rs and gui.py. This manual describes version 2.0.0 as inspected on 14 October 2026.')
out=ROOT/'docs/DAS_Transcriber_User_Manual.docx';d.save(out)
print(out)
print('Paragraphs:',len(d.paragraphs),'Tables:',len(d.tables),'Images:',len(d.inline_shapes))
