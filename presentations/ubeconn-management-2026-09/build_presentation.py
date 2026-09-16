"""Build an editable executive deck and a matching, font-stable PDF.

Run from repository root with Python 3 + python-pptx, Pillow, fonttools[brotli],
matplotlib and reportlab. No source data, application or database is modified.
"""
from pathlib import Path
import csv
import json
import math
import shutil
import textwrap
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont
from fontTools.ttLib import TTFont
from fontTools.pens.svgPathPen import SVGPathPen
from fontTools.pens.transformPen import TransformPen
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont as PDFTTFont

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
ASSETS = OUT / 'assets'
FONTS = OUT / 'fonts'
PREVIEWS = OUT / 'previews'
for d in (ASSETS, FONTS, PREVIEWS):
    d.mkdir(exist_ok=True)

C = dict(bg='F4EFFA', panel='FFFFFF', line='D7C9E8', purple='7548BA',
         lilac='6D3CAE', white='FFFFFF', fg='2E1E43', dim='6B5B7D', pink='9764C3')
FONT_SPECS = {
    'body': ('ibm-plex-sans', 400, 'IBM Plex Sans'),
    'medium': ('ibm-plex-sans', 700, 'IBM Plex Sans'),
    'head': ('ibm-plex-sans-condensed', 700, 'IBM Plex Sans Condensed'),
    'mono': ('ibm-plex-mono', 400, 'IBM Plex Mono'),
}
FONT_PATHS = {}
for role, (family, weight, name) in FONT_SPECS.items():
    source = ROOT / f'web/node_modules/@fontsource/{family}/files/{family}-latin-{weight}-normal.woff2'
    dest = FONTS / f'{family}-{weight}.ttf'
    font = TTFont(source)
    font.flavor = None
    font.save(dest)
    FONT_PATHS[role] = dest
    pdfmetrics.registerFont(PDFTTFont(role, str(dest)))
shutil.copyfile(ROOT / 'web/node_modules/@fontsource/ibm-plex-sans/LICENSE', FONTS / 'LICENSE.txt')

# Supplied logo is embedded unchanged, including its transparency.
LOGO = ASSETS / 'ubeconn-supplied-logo.png'

MASTS = list(csv.DictReader((ROOT / 'results/masts/masts.csv').open()))
SUMMARY = json.loads((ROOT / 'results/das-discovery-pilot-20260916/summary.json').read_text())
BENCH = json.loads((ROOT / 'results/das-discovery-pilot-20260916/pilot-validation.json').read_text())
assert len(MASTS) == 579
assert SUMMARY['feature_rows'] == 240

W, H = 13.333333, 7.5
SLIDES = []
CURRENT = None

def obj(kind, **kw):
    CURRENT['objects'].append(dict(kind=kind, **kw))

def rect(x, y, w, h, color='panel'):
    obj('rect', x=x, y=y, w=w, h=h, color=C.get(color, color))

def line(x1, y1, x2, y2, color='line', width=1):
    obj('line', x=x1, y=y1, x2=x2, y2=y2, color=C.get(color, color), width=width)

def text(s, x, y, w, size=20, role='body', color='fg', h=None):
    # Explicit line breaks are shared by all output formats.
    s = str(s).replace('→', '>').replace('≈', '~').replace('≠', 'is not')
    f = ImageFont.truetype(str(FONT_PATHS[role]), int(size * 4))
    lines = []
    limit = w * 72 * 4 - 4
    for paragraph in str(s).split('\n'):
        words = paragraph.split(' ')
        current = ''
        for word in words:
            candidate = (current + ' ' + word).strip()
            if current and f.getlength(candidate) > limit:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    height = h or (len(lines) * size * 1.22 / 72 + 0.07)
    obj('text', text='\n'.join(lines), x=x, y=y, w=w, h=height, size=size,
        role=role, color=C.get(color, color))
    return height

def image(path, x, y, w=None, h=None):
    im = Image.open(path)
    if w is None:
        w = h * im.width / im.height
    if h is None:
        h = w * im.height / im.width
    obj('image', path=str(path), x=x, y=y, w=w, h=h)

def slide(kicker, title, subtitle, source, notes=''):
    global CURRENT
    CURRENT = dict(objects=[], title=title, source=source, notes=notes)
    SLIDES.append(CURRENT)
    text(kicker.upper(), .55, .30, 10.4, 10, 'mono', 'lilac')
    image(LOGO, 10.61, 0, w=2.35)
    text(title, .55, .83, 12.1, 33, 'head')
    text(subtitle, .58, 1.48, 12.0, 14, 'body', 'dim')
    line(.55, 1.98, 12.78, 1.98, 'lilac', 1.1)
    line(.55, 7.02, 12.78, 7.02)
    text(source, .57, 7.14, 10.95, 8, 'mono', 'dim')
    text(f'UBE / {len(SLIDES):02d}', 11.67, 7.13, 1.10, 9, 'mono', 'lilac')

def takeaway(s):
    rect(.55, 6.46, 12.23, .39)
    rect(.55, 6.46, .035, .39, 'purple')
    text(s, .72, 6.52, 11.88, 13, 'medium', 'lilac')

def card(x, y, w, h, label, title, body, title_size=26):
    rect(x, y, w, h)
    rect(x, y, w, .035, 'purple')
    text(label.upper(), x + .20, y + .20, w - .40, 10, 'mono', 'lilac')
    th = text(title, x + .20, y + .57, w - .40, title_size, 'head')
    text(body, x + .20, y + .68 + th, w - .40, 17, 'body', 'dim')

def stat(x, y, w, value, label):
    text(value, x, y, w, 43, 'head', 'lilac')
    text(label, x, y + .79, w, 13, 'body', 'dim')

def route_plot(x, y, w, h, cover=False):
    colors = ['CFADF0', '9D6BD0', 'DFAFD2', 'F6F0FC']
    xs = [float(r['longitude']) * math.cos(math.radians(68.15)) for r in MASTS]
    ys = [float(r['latitude']) for r in MASTS]
    scale = min(w / (max(xs) - min(xs)), h / (max(ys) - min(ys)))
    ox = x + (w - (max(xs) - min(xs)) * scale) / 2
    oy = y + (h - (max(ys) - min(ys)) * scale) / 2
    for route, color in zip(['R06', 'R12', 'R13', 'R14'], colors):
        rows = sorted([r for r in MASTS if r['route'] == route], key=lambda r: int(r['route_position']))
        points = [(ox + (float(r['longitude']) * math.cos(math.radians(68.15)) - min(xs)) * scale,
                   oy + (max(ys) - float(r['latitude'])) * scale) for r in rows]
        for a, b in zip(points, points[1:]):
            line(*a, *b, color, 2.0 if cover else 1.7)
        for px, py in points[::8]:
            obj('circle', x=px-.025, y=py-.025, w=.05, h=.05, color=color)


# 01 / A management story, with the supplied logo unchanged.
slide('Ubeconn AS / Management update', 'Moving forward with event classification',
      'Using machine learning to classify real cable events, with verified examples and expert review.',
      'PROJECT UPDATE / 16 SEPTEMBER 2026',
      'Prepared for executive management. This is a new, less technical presentation, not an update to the previous deck. It uses the supplied purple logo unchanged and a brighter ube palette. Company name remains Ubeconn AS in text; the supplied artwork is not respelled. Current evidence supports a local investigation platform, a completed discovery pilot and supervised training software. It does not establish a field-validated classifier. Weather migration 005 is deployed on the server and local installation and supersedes the older two-location/no-humidity snapshot.')
image(LOGO, .43, 2.00, w=7.32)
text('A clearer view of cable behaviour.', .85, 4.92, 6.56, 28, 'head')
text('Better context. Structured review.\nA foundation for learning.', .87, 5.55, 6.50, 20, 'body', 'dim')
card(8.04, 2.30, 4.74, 3.71, 'Our focus', 'Classify real\ncable events',
     'Bring recordings, weather and expert knowledge together.\n\nUse confirmed examples to teach the model and test its suggestions.', 29)
takeaway('The platform is ready for structured investigation. Verified examples are the next priority.')

# 02
slide('01 / Platform foundation', 'The building blocks for event classification are in place',
      'The next phase brings these capabilities together around real recordings and confirmed examples.',
      'EVIDENCE / project documentation; detailed sources and status in speaker notes',
      'Sources: docs/map-app.md; docs/masts.md; docs/weather.md; docs/event-ml.md. User clarification: existing event entries were randomly generated demonstration data, not a historical incident archive or candidate training set. They are intended for future deletion; this presentation task does not delete database records. The local application supports investigation and review. The asset import contains 579 masts across four routes and 575 inferred links. Weather collection is deployed. Training code exists; a field-validated classifier still requires confirmed examples and testing.')
text('CAPABILITY', .77, 2.27, 4.00, 11, 'mono', 'lilac')
text('WHAT IT ENABLES', 5.31, 2.27, 7.10, 11, 'mono', 'lilac')
rows=[('Unified workspace', 'Explore recordings and retain expert conclusions'),
      ('Mapped assets', 'Investigate cable behaviour across four routes'),
      ('17 weather locations', 'Add temperature, wind and humidity context'),
      ('Discovery workflow', 'Find useful examples for expert review'),
      ('Machine learning', 'Random Forest classification trained on verified examples')]
for i,(a,b) in enumerate(rows):
    yy=2.77+i*.64
    if i%2==0:rect(.55, yy-.04, 12.23, .64)
    text(a,.77,yy+.07,4.18,21,'head')
    text(b,5.31,yy+.09,7.15,18,'body','dim')
takeaway('Our focus is to turn real recordings and expert knowledge into dependable event classifications.')

# 03
slide('02 / The operator experience', 'A clear workflow for reviewing real recordings',
      'The local workspace supports investigation, expert review and a lasting record of the evidence.',
      'WORKFLOW / based on the implemented local investigation application',
      'Source: docs/map-app.md. The application supports asset search, signal inspection, archive replay and classification reviews with evidence notes and history. The prior screenshot was removed because it displayed randomly generated demonstration events. No demonstration event counts are presented as operational evidence. Classifying real recordings requires explicit verified time and location references. This workflow diagram describes capabilities, not a claim of production rollout.')
card(.55,2.30,3.94,3.22,'01 / Explore','Choose a real recording',
     'Find the relevant asset and inspect the recorded cable behaviour.',27)
card(4.69,2.30,3.94,3.22,'02 / Understand','Bring the evidence together',
     'Compare signal patterns with weather, inspections and confirmed observations.',27)
card(8.83,2.30,3.95,3.22,'03 / Record','Save the reviewed category',
     'Keep the conclusion, reviewer and supporting evidence for learning and follow-up.',27)
text('RECORDING + CONTEXT + EXPERT REVIEW',.78,5.93,11.80,16,'mono','lilac')
takeaway('Every training example should connect a real recording to an evidence-backed label.')

# UI screenshot / added after the operator workflow
slide('03 / The new user interface', 'The new UI brings map, signals and review together',
      'A single workspace helps reviewers move from cable behaviour to a documented conclusion.',
      'SCREENSHOT / local application; demonstration entries shown for interface illustration only',
      'Sources: docs/map-app.md and assets/application-overview.png. Screenshot is an unchanged capture of the local application. Event entries and counts visible in the screen are randomly generated demonstration data, as clarified by the user; they are not real incidents, a training set or performance evidence. The slide illustrates interface capabilities only. The map displays assets and signal layers; the recording browser supports waveform, heatmap, spectrum and playback. The review panel supports a category, reviewer, evidence note and review history when an event is selected. No event is selected in this overview screenshot. The displayed recording is an independent archive example. Screenshot retains the original app branding and colours, and basemap attribution: OpenFreeMap / OpenMapTiles / OpenStreetMap. The supplied Ubeconn logo and pale ube slide palette remain unchanged.')
image(ASSETS/'application-overview.png', .58, 2.15, h=4.50)
for yy,title,body in [
    (2.25,'1. Explore the network','Search masts and routes. Use the map to choose where to investigate.'),
    (3.58,'2. Inspect the recording','Review signal patterns and play back the recording to see changes over time.'),
    (4.91,'3. Record the conclusion','Select an event to save its category, reviewer and supporting evidence.')]:
    rect(5.89,yy,6.89,1.16)
    text(title,6.11,yy+.13,6.45,24,'head')
    text(body,6.12,yy+.63,6.43,16,'body','dim')
text('UI DEMONSTRATION',6.11,6.29,6.4,10,'mono','lilac')
text('Visible event entries are illustrative, not real incidents.',6.11,6.61,6.4,13,'body','dim')

# 04
slide('04 / Weather upgrade — deployed', 'Weather context now follows more of the network',
      'Forecast coverage has expanded, and relative humidity is now being collected.',
      'EVIDENCE / docs/weather.md · migration 005 applied on server and local installation',
      'Migration 005_distributed_weather.sql and the collector were deployed on 16 September 2026. Two original locations remain; fifteen new points use mast coordinates across R06, R12, R13 and R14 at approximately five kilometre spacing. There are now 17 locations × four weather variables. Both installations recorded successful initial captures of 68 rows. Server capture runs hourly at minute 13; local capture is on demand. These are MET/YR modeled forecasts, not direct measurements. Source validity, issue and retrieval timestamps are stored on new rows. Integration with ML requires verified recording times, appropriate location selection and weather available at the time of the event.')
rect(.55,2.25,5.29,2.02)
text('BEFORE',.79,2.48,4.65,11,'mono','dim')
text('2',.79,2.91,1.05,56,'head','dim')
text('forecast locations',2.01,3.12,3.49,25,'head','dim')
rect(6.20,2.25,6.58,2.02)
rect(6.20,2.25,.045,2.02,'purple')
text('NOW',6.49,2.48,5.91,11,'mono','lilac')
text('17',6.49,2.91,1.42,56,'head','lilac')
text('locations across the route network',8.08,3.04,4.36,24,'head')
for i,(name,status) in enumerate([('Temperature','COLLECTED'),('Wind speed','COLLECTED'),
                                ('Wind direction','COLLECTED'),('Humidity','NEW')]):
    xx=.55+i*3.11
    rect(xx,4.62,2.90,1.16,'purple' if i==3 else 'panel')
    text(status,xx+.20,4.83,2.49,10,'mono','white' if i==3 else 'lilac')
    text(name,xx+.20,5.18,2.49,23,'head','white' if i==3 else 'fg')
text('New records also retain when the forecast applies and when it was issued.',.78,6.00,11.9,17,'body','dim')
takeaway('Richer weather context supports event interpretation; forecasts do not confirm the cause of an event.')

# 05
slide('05 / Machine learning, in plain language', 'We use Random Forest machine learning for classification',
      'Supervised learning: confirmed examples teach the model to recognise the four event categories.',
      'MODEL STATUS / supervised training software implemented; field accuracy not yet established',
      'Source: src/Capture/event_ml.py and docs/event-ml.md. The supervised baseline is a Random Forest with 300 decision trees and balanced class weights. Training uses signal summaries covering vibration strength, peaks, rhythmic behaviour, variability across channels and changes over the window, plus optional temperature, wind speed/direction and humidity. Missing weather is explicit. The model learns rules from reviewed examples; a new window receives a predicted class, four uncalibrated scores and a low-score review flag. No supervised trained model or measured field accuracy is supplied. The scores must not be presented as a validated confidence percentage. Weather must be aligned before being used. The baseline outputs a single category for each window, not an incident duration or severity.')
card(.55,2.26,3.94,2.80,'The clues','Cable behaviour + weather',
     'How strong is the vibration?\nIs it rhythmic or irregular?\nWhat were the conditions?',25)
card(4.69,2.26,3.94,2.80,'Machine learning technique','Random Forest',
     '300 decision trees learn from confirmed examples. Their combined votes suggest the event category.',25)
card(8.83,2.26,3.95,2.80,'The result','A suggested category for review',
     'The output includes a class and scores. Uncertain cases are flagged for attention.',25)
text('INTENDED OUTPUTS',.76,5.35,11.85,10,'mono','lilac')
for i,name in enumerate(['Normal operations','Ice loading','Wind galloping','External interference']):
    xx=.55+i*3.11
    rect(xx,5.76,2.90,.46)
    text(name,xx+.14,5.84,2.64,17,'head','lilac')
takeaway('The software can learn and make predictions; its reliability must be demonstrated on independent events.')

# 06
slide('06 / The learning process', 'Two machine-learning steps support event classification',
      'Unsupervised learning finds patterns; supervised learning links those patterns to confirmed event types.',
      'EVIDENCE / docs/event-discovery.md · recorded discovery pilot · docs/event-ml.md',
      'The offline discovery workflow uses clustering (KMeans) and anomaly ranking (Isolation Forest), without verified class labels. The recorded pilot covered roughly 15 minutes across 90 archive files, producing 240 window/block feature rows and six exploratory clusters. These clusters are not the four event classes. No weather join was used in that pilot. Human reviewers must use independent evidence, including inspections, cameras or confirmed incident records, to assign labels. A separate supervised Random Forest can then train on those examples. Related windows from one session/incident must stay in the same evaluation group, with additional unseen sites/seasons/incidents reserved for final assessment. Neither the pilot nor its six clusters demonstrates event-classification accuracy.')
steps=[('01 / PILOT COMPLETED','Find patterns',
        'K-means groups similar signals. Isolation Forest highlights unusual ones.'),
       ('02 / NEXT PRIORITY','Confirm the cause',
        'People use inspections or incident evidence to assign reliable labels.'),
       ('03 / SOFTWARE READY','Teach the model',
        'Random Forest learns the four event categories from expert-labeled examples.'),
       ('04 / BEFORE ROLLOUT','Test new events',
        'Check its suggestions on events it has never seen during training.')]
for i,(label,title,body) in enumerate(steps):
    card(.55+i*3.11,2.28,2.90,3.24,label,title,body,27)
text('PILOT EVIDENCE',.77,5.90,2.91,11,'mono','lilac')
text('Discovery exercised on about 15 minutes of real recordings.',3.85,5.78,8.61,23,'head')
takeaway('Finding an unusual pattern is a way to select examples for review; it does not establish an event class.')

# 07
slide('07 / The four categories', 'The goal is to distinguish four types of operation',
      'Each category needs confirmed examples, including cases that are easy to confuse.',
      'TARGET CATEGORIES / classification design; detection performance still to be validated',
      'Source: docs/event-ml.md; project objectives in AGENTS.md. Exact database class names are Normal Operations, Ice Loading, Wind Galloping and External Interference. Vegetation contact or strikes are examples for the External Interference category, not an independently validated vegetation classifier. Normal examples must include quiet and windy operation. Ice detection may need additional evidence, sensors or features; existing vibration features alone may not distinguish it reliably. The current model chooses one class per window and cannot express simultaneous ice and galloping. Class names do not establish severity, ice mass or loading force.')
classes=[('01','Normal operations','Recognise normal cable behaviour, including windy conditions.'),
         ('02','Ice loading','Test for changes associated with independently confirmed ice.'),
         ('03','Wind galloping','Recognise sustained cable movement confirmed in the field.'),
         ('04','External interference','Investigate disturbances such as strikes or vegetation contact.')]
for i,(label,title,body) in enumerate(classes):
    xx=.55+(i%2)*6.23
    yy=2.27+(i//2)*1.84
    rect(xx,yy,6.00,1.57)
    text(label,xx+.21,yy+.22,.66,12,'mono','lilac')
    text(title,xx+1.02,yy+.19,4.72,27,'head')
    text(body,xx+1.02,yy+.78,4.68,17,'body','dim')
takeaway('A quiet signal is not automatically normal, and cold or humid weather is not proof of ice.')

# 08
slide('08 / Build the training dataset', 'Collect confirmed examples of real cable behaviour',
      'Build a dedicated training dataset from verified recordings and independent field evidence.',
      'NEXT PHASE / real recordings, expert labels and independent evaluation',
      'Sources: user clarification; docs/event-ml.md; docs/map-app.md. Existing event records are randomly generated demonstration data and are excluded from the classification dataset. The user intends to delete them later; no database deletion is performed in this task. Collect verified real recordings for all four target classes, including varied normal operation. Retain recording/time/location references and independent evidence for every label. Keep related windows from the same incident grouped when evaluating the model. No amount of synthetic demonstration events establishes field performance.')
rect(.55,2.28,4.04,3.75)
text('Real data',.81,2.83,3.57,40,'head','lilac')
text('Verified labels',.83,3.78,3.52,28,'head')
text('A dedicated dataset for training and testing the classifier.',.83,4.48,3.49,20,'body','dim')
items=[('Match the evidence','Confirm when and where a recording belongs.'),
       ('Verify the cause','Use inspections, cameras or reliable incident reports.'),
       ('Include normal operation','Teach the model what ordinary conditions look like.')]
for i,(title,body) in enumerate(items):
    yy=2.38+i*1.18
    text(f'0{i+1}',5.03,yy+.07,.67,13,'mono','lilac')
    text(title,5.87,yy,6.66,26,'head')
    text(body,5.88,yy+.51,6.63,18,'body','dim')
takeaway('The most valuable next asset is a small, trustworthy set of examples that can grow over time.')

# 09
slide('09 / Next steps', 'Start with verified examples, then classify with oversight',
      'A short sequence takes the project from the current platform to useful event suggestions.',
      'PROPOSED NEXT STEPS / scope, owners and acceptance targets to be agreed',
      'Proposed sequence grounded in docs/event-ml.md, docs/map-app.md and docs/weather.md. 1: select a manageable pilot scope and verify recording time, channel/location mapping, preprocessing and weather availability. 2: field experts label representative examples for all four classes, keeping independent incidents grouped and retaining evidence. 3: train the Random Forest and assess on independent events, measuring per-class precision/recall, false alerts, missed events and practical reviewer workload; confidence scores and thresholds need validation. 4: run class suggestions in a controlled pilot with human review before automating operational escalation. Model output is not yet integrated as a proven production classifier. No budget, schedule, numerical performance or customer commitment is implied. Initial management action: name a field-review owner and agree the pilot scope.')
steps=[('1','Select and align','Choose pilot recordings and verify time, place and weather.'),
       ('2','Confirm the labels','Have experts review examples of all four categories.'),
       ('3','Train and test','Measure correct classifications, missed events and false alerts.'),
       ('4','Use with review','Pilot model suggestions while people check the results.')]
for i,(n,title,body) in enumerate(steps):
    xx=.55+i*3.11
    card(xx,2.30,2.90,3.15,'STEP '+n,title,body,26)
rect(.55,5.73,12.23,.47,'purple')
text('FIRST MILESTONE: an expert-reviewed dataset with evidence for every label.',.77,5.83,11.77,17,'medium','white')
takeaway('Management focus: agree the pilot scope and assign a field-review owner to build the first labeled dataset.')

def rgb(s):
    return tuple(int(s[i:i+2],16) for i in (0,2,4))

def write_pptx():
    prs=Presentation()
    prs.slide_width=Inches(W)
    prs.slide_height=Inches(H)
    prs.core_properties.title='Ubeconn AS | Cable intelligence — executive project review'
    prs.core_properties.subject='Project progress, technical evidence and proposed validation roadmap'
    prs.core_properties.author='Ubeconn AS'
    prs.core_properties.keywords='Ubeconn, DAS, cable monitoring, executive review, September 2026'
    prs.core_properties.comments='Editable slide content. Source details and caveats are in speaker notes.'
    for d in SLIDES:
        s=prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb=RGBColor.from_string(C['bg'])
        for o in d['objects']:
            kind=o['kind']
            if kind in ('rect','circle'):
                shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE if kind=='rect' else MSO_SHAPE.OVAL,
                                        Inches(o['x']),Inches(o['y']),Inches(o['w']),Inches(o['h']))
                shape.fill.solid()
                shape.fill.fore_color.rgb=RGBColor.from_string(o['color'])
                shape.line.fill.background()
            elif kind=='line':
                shape=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(o['x']),Inches(o['y']),
                                            Inches(o['x2']),Inches(o['y2']))
                shape.line.color.rgb=RGBColor.from_string(o['color'])
                shape.line.width=Pt(o['width'])
            elif kind=='image':
                s.shapes.add_picture(o['path'],Inches(o['x']),Inches(o['y']),Inches(o['w']),Inches(o['h']))
            elif kind=='text':
                shape=s.shapes.add_textbox(Inches(o['x']),Inches(o['y']),Inches(o['w']),Inches(o['h']))
                tf=shape.text_frame
                tf.clear()
                tf.word_wrap=False
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                for j,txt in enumerate(o['text'].split('\n')):
                    p=tf.paragraphs[0] if j==0 else tf.add_paragraph()
                    p.text=txt
                    p.space_before=Pt(0)
                    p.space_after=Pt(0)
                    p.line_spacing=Pt(o['size']*1.22)
                    p.font.name=FONT_SPECS[o['role']][2]
                    p.font.size=Pt(o['size'])
                    p.font.bold=o['role'] in ('head','medium')
                    p.font.color.rgb=RGBColor.from_string(o['color'])
        s.notes_slide.notes_text_frame.text=d['notes']+'\n\nSources: '+d['source']
    # Theme: Ubeconn colours and IBM Plex, so new editable objects inherit the brand.
    theme=prs.slide_masters[0].part.theme_part if hasattr(prs.slide_masters[0].part,'theme_part') else None
    from lxml import etree
    for part in prs.part.package.iter_parts():
        if str(part.partname).startswith('/ppt/theme/theme'):
            tree=etree.fromstring(part.blob)
            ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
            clr=tree.find('.//a:clrScheme',ns)
            clr.set('name','Ubeconn')
            scheme=['2E1E43','F4EFFA','6B5B7D','FFFFFF','7548BA','6D3CAE','9764C3','D7C9E8','BAA2DC','EAE0F4','6D3CAE','9764C3']
            for node,color in zip(clr,scheme):
                for child in list(node):node.remove(child)
                ch=etree.SubElement(node,'{'+ns['a']+'}srgbClr');ch.set('val',color)
            tree.find('.//a:majorFont/a:latin',ns).set('typeface','IBM Plex Sans Condensed')
            tree.find('.//a:minorFont/a:latin',ns).set('typeface','IBM Plex Sans')
            part._blob=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    prs.save(OUT/'Ubeconn_Management_Update_2026-09.pptx')


def write_previews_and_pdf():
    scale=144 # 1920 × 1080 preview
    pdf=canvas.Canvas(str(OUT/'Ubeconn_Management_Update_2026-09.pdf'),pagesize=(W*72,H*72))
    pdf.setTitle('Ubeconn AS | Management project update | September 2026')
    pdf.setAuthor('Ubeconn AS')
    for index,d in enumerate(SLIDES,1):
        im=Image.new('RGB',(round(W*scale),round(H*scale)),rgb(C['bg']))
        draw=ImageDraw.Draw(im)
        pdf.setFillColorRGB(*(v/255 for v in rgb(C['bg'])))
        pdf.rect(0,0,W*72,H*72,fill=1,stroke=0)
        for o in d['objects']:
            k=o['kind'];x=o['x'];y=o['y']
            if k in ('rect','circle'):
                box=(round(x*scale),round(y*scale),round((x+o['w'])*scale),round((y+o['h'])*scale))
                pdf.setFillColorRGB(*(v/255 for v in rgb(o['color'])))
                if k=='rect':
                    draw.rectangle(box,fill=rgb(o['color']))
                    pdf.rect(x*72,(H-y-o['h'])*72,o['w']*72,o['h']*72,fill=1,stroke=0)
                else:
                    draw.ellipse(box,fill=rgb(o['color']))
                    pdf.ellipse(x*72,(H-y-o['h'])*72,(x+o['w'])*72,(H-y)*72,fill=1,stroke=0)
            elif k=='line':
                draw.line((round(x*scale),round(y*scale),round(o['x2']*scale),round(o['y2']*scale)),
                          fill=rgb(o['color']),width=max(1,round(o['width']*scale/72)))
                pdf.setStrokeColorRGB(*(v/255 for v in rgb(o['color'])))
                pdf.setLineWidth(o['width'])
                pdf.line(x*72,(H-y)*72,o['x2']*72,(H-o['y2'])*72)
            elif k=='image':
                pic=Image.open(o['path']).convert('RGBA').resize((round(o['w']*scale),round(o['h']*scale)),Image.Resampling.LANCZOS)
                im.paste(pic,(round(x*scale),round(y*scale)),pic)
                pdf.drawImage(o['path'],x*72,(H-y-o['h'])*72,o['w']*72,o['h']*72,mask='auto')
            elif k=='text':
                f=ImageFont.truetype(str(FONT_PATHS[o['role']]),round(o['size']*scale/72))
                pdf.setFont(o['role'],o['size'])
                pdf.setFillColorRGB(*(v/255 for v in rgb(o['color'])))
                for j,t in enumerate(o['text'].split('\n')):
                    baseline=y*scale+(o['size']*.97+j*o['size']*1.22)*scale/72
                    draw.text((x*scale,baseline),t,font=f,anchor='ls',fill=rgb(o['color']))
                    pdf.drawString(x*72,(H-y)*72-o['size']*.97-j*o['size']*1.22,t)
        im.save(PREVIEWS/f'slide-{index:02d}.png')
        pdf.showPage()
    pdf.save()
    thumbw=480;thumbh=270
    sheet=Image.new('RGB',(thumbw*3,thumbh*math.ceil(len(SLIDES)/3)),(15,10,21))
    for i in range(len(SLIDES)):
        im=Image.open(PREVIEWS/f'slide-{i+1:02d}.png').resize((thumbw,thumbh),Image.Resampling.LANCZOS)
        sheet.paste(im,((i%3)*thumbw,(i//3)*thumbh))
    sheet.save(OUT/'slide-overview.png')


def validate():
    issues=[]
    for i,d in enumerate(SLIDES,1):
        for o in d['objects']:
            if o['kind']=='line':continue
            if o['x'] < -.01 or o['y'] < -.01 or o['x']+o['w'] > W+.01 or o['y']+o['h'] > H+.01:
                issues.append((i,'outside slide',o.get('text','')))
            if o['kind']=='text':
                for txt in o['text'].split('\n'):
                    font=ImageFont.truetype(str(FONT_PATHS[o['role']]),round(o['size']*4))
                    if font.getlength(txt) > o['w']*72*4:
                        issues.append((i,'text overflow',txt))
                # Header and source text excepted: all body text must stop above the footer.
                if 2.05<o['y']<6.45 and o['y']+o['h']>6.95:
                    issues.append((i,'body reaches footer',o['text']))
    if issues:raise RuntimeError(issues)
    p=Presentation(OUT/'Ubeconn_Management_Update_2026-09.pptx')
    assert len(p.slides)==10
    assert all(s.has_notes_slide for s in p.slides)
    (OUT/'validation.json').write_text(json.dumps({'slides':len(p.slides),'bounds_and_text_width':'passed',
        'speaker_notes':'all slides','preview_method':'shared layout rendered with bundled fonts; not an Office-engine render',
        'source_masts':len(MASTS),'pilot_feature_rows':SUMMARY['feature_rows']},indent=2))

write_pptx()
validate()
write_previews_and_pdf()
(OUT/'speaker-notes.md').write_text('\n\n'.join(f'## {i:02d} — {d["title"]}\n\n{d["notes"]}\n\n{d["source"]}' for i,d in enumerate(SLIDES,1)))
(OUT/'brand-palette.json').write_text(json.dumps(C,indent=2))
print(f'Wrote {len(SLIDES)} slides to {OUT}')
