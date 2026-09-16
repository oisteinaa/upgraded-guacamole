"""Build an editable executive deck and a matching, font-stable PDF.

Run from repository root with Python 3 + python-pptx, Pillow, fonttools[brotli],
matplotlib and reportlab. No source data, application or database is modified.
"""
from pathlib import Path
import csv
import json
import math
import shutil
import textwrap
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont
from fontTools.ttLib import TTFont
from fontTools.pens.svgPathPen import SVGPathPen
from fontTools.pens.transformPen import TransformPen
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont as PDFTTFont

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
ASSETS = OUT / 'assets'
FONTS = OUT / 'fonts'
PREVIEWS = OUT / 'previews'
for d in (ASSETS, FONTS, PREVIEWS):
    d.mkdir(exist_ok=True)

C = dict(bg='1D1429', panel='2B1E3A', line='493458', purple='9D6BD0',
         lilac='CFADF0', white='FFFFFF', fg='F6F0FC', dim='BAA9CB', pink='DFAFD2')
FONT_SPECS = {
    'body': ('ibm-plex-sans', 400, 'IBM Plex Sans'),
    'medium': ('ibm-plex-sans', 700, 'IBM Plex Sans'),
    'head': ('ibm-plex-sans-condensed', 700, 'IBM Plex Sans Condensed'),
    'mono': ('ibm-plex-mono', 400, 'IBM Plex Mono'),
}
FONT_PATHS = {}
for role, (family, weight, name) in FONT_SPECS.items():
    source = ROOT / f'web/node_modules/@fontsource/{family}/files/{family}-latin-{weight}-normal.woff2'
    dest = FONTS / f'{family}-{weight}.ttf'
    font = TTFont(source)
    font.flavor = None
    font.save(dest)
    FONT_PATHS[role] = dest
    pdfmetrics.registerFont(PDFTTFont(role, str(dest)))
shutil.copyfile(ROOT / 'web/node_modules/@fontsource/ibm-plex-sans/LICENSE', FONTS / 'LICENSE.txt')

# A plain white wordmark, saved as vector outlines for reliable reuse.
logo_font = TTFont(FONT_PATHS['medium'])
glyphs = logo_font.getGlyphSet()
cmap = logo_font.getBestCmap()
pen = SVGPathPen(glyphs)
advance = 0
for character in 'ubeconn':
    name = cmap[ord(character)]
    glyphs[name].draw(TransformPen(pen, (1, 0, 0, -1, advance, 850)))
    advance += logo_font['hmtx'][name][0]
(ASSETS / 'ubeconn-white.svg').write_text(
    f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {advance} 1100">'
    f'<title>Ubeconn white wordmark</title><path fill="#fff" d="{pen.getCommands()}"/></svg>')
logo = Image.new('RGBA', (advance, 1100), (0, 0, 0, 0))
ImageDraw.Draw(logo).text((0, 850), 'ubeconn', font=ImageFont.truetype(str(FONT_PATHS['medium']), 1000),
                         anchor='ls', fill='white')
logo.save(ASSETS / 'ubeconn-white.png')

MASTS = list(csv.DictReader((ROOT / 'results/masts/masts.csv').open()))
SUMMARY = json.loads((ROOT / 'results/das-discovery-pilot-20260916/summary.json').read_text())
BENCH = json.loads((ROOT / 'results/das-discovery-pilot-20260916/pilot-validation.json').read_text())
assert len(MASTS) == 579
assert SUMMARY['feature_rows'] == 240

W, H = 13.333333, 7.5
SLIDES = []
CURRENT = None

def obj(kind, **kw):
    CURRENT['objects'].append(dict(kind=kind, **kw))

def rect(x, y, w, h, color='panel'):
    obj('rect', x=x, y=y, w=w, h=h, color=C.get(color, color))

def line(x1, y1, x2, y2, color='line', width=1):
    obj('line', x=x1, y=y1, x2=x2, y2=y2, color=C.get(color, color), width=width)

def text(s, x, y, w, size=20, role='body', color='fg', h=None):
    # Explicit line breaks are shared by all output formats.
    s = str(s).replace('→', '>').replace('≈', '~').replace('≠', 'is not')
    f = ImageFont.truetype(str(FONT_PATHS[role]), int(size * 4))
    lines = []
    limit = w * 72 * 4 - 4
    for paragraph in str(s).split('\n'):
        words = paragraph.split(' ')
        current = ''
        for word in words:
            candidate = (current + ' ' + word).strip()
            if current and f.getlength(candidate) > limit:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    height = h or (len(lines) * size * 1.22 / 72 + 0.07)
    obj('text', text='\n'.join(lines), x=x, y=y, w=w, h=height, size=size,
        role=role, color=C.get(color, color))
    return height

def image(path, x, y, w=None, h=None):
    im = Image.open(path)
    if w is None:
        w = h * im.width / im.height
    if h is None:
        h = w * im.height / im.width
    obj('image', path=str(path), x=x, y=y, w=w, h=h)

def slide(kicker, title, subtitle, source, notes=''):
    global CURRENT
    CURRENT = dict(objects=[], title=title, source=source, notes=notes)
    SLIDES.append(CURRENT)
    text(kicker.upper(), .55, .30, 10.4, 10, 'mono', 'lilac')
    text('ubeconn', 11.42, .25, 1.40, 23, 'medium', 'white')
    text(title, .55, .83, 12.1, 33, 'head')
    text(subtitle, .58, 1.48, 12.0, 14, 'body', 'dim')
    line(.55, 1.98, 12.78, 1.98, 'lilac', 1.1)
    line(.55, 7.02, 12.78, 7.02)
    text(source, .57, 7.14, 10.95, 8, 'mono', 'dim')
    text(f'UBE / {len(SLIDES):02d}', 11.67, 7.13, 1.10, 9, 'mono', 'lilac')

def takeaway(s):
    rect(.55, 6.46, 12.23, .39)
    rect(.55, 6.46, .035, .39, 'purple')
    text(s, .72, 6.52, 11.88, 13, 'medium', 'lilac')

def card(x, y, w, h, label, title, body, title_size=26):
    rect(x, y, w, h)
    rect(x, y, w, .035, 'purple')
    text(label.upper(), x + .20, y + .20, w - .40, 10, 'mono', 'lilac')
    th = text(title, x + .20, y + .57, w - .40, title_size, 'head')
    text(body, x + .20, y + .68 + th, w - .40, 17, 'body', 'dim')

def stat(x, y, w, value, label):
    text(value, x, y, w, 43, 'head', 'lilac')
    text(label, x, y + .79, w, 13, 'body', 'dim')

def route_plot(x, y, w, h, cover=False):
    colors = ['CFADF0', '9D6BD0', 'DFAFD2', 'F6F0FC']
    xs = [float(r['longitude']) * math.cos(math.radians(68.15)) for r in MASTS]
    ys = [float(r['latitude']) for r in MASTS]
    scale = min(w / (max(xs) - min(xs)), h / (max(ys) - min(ys)))
    ox = x + (w - (max(xs) - min(xs)) * scale) / 2
    oy = y + (h - (max(ys) - min(ys)) * scale) / 2
    for route, color in zip(['R06', 'R12', 'R13', 'R14'], colors):
        rows = sorted([r for r in MASTS if r['route'] == route], key=lambda r: int(r['route_position']))
        points = [(ox + (float(r['longitude']) * math.cos(math.radians(68.15)) - min(xs)) * scale,
                   oy + (max(ys) - float(r['latitude'])) * scale) for r in rows]
        for a, b in zip(points, points[1:]):
            line(*a, *b, color, 2.0 if cover else 1.7)
        for px, py in points[::8]:
            obj('circle', x=px-.025, y=py-.025, w=.05, h=.05, color=color)

# 01 — cover
slide('Ubeconn AS / Executive project review', 'From fibre signals to cable intelligence',
      'Distributed acoustic sensing for cable monitoring • Project status and proposed next steps',
      'PROJECT SNAPSHOT / 16 SEP 2026',
      'Audience: technically informed executives and project stakeholders. This presentation describes the repository and recorded project evidence as of 16 September 2026. It is not a claim of a production rollout or certified event detection. Future phases are proposed, with no approved budget or delivery dates implied.')
rect(.55, 2.23, 6.25, 3.95)
text('SEE THE SIGNAL.\nUNDERSTAND THE EVENT.', .85, 2.63, 5.70, 32, 'head')
text('A traceable path from DAS recordings to mapped evidence, expert review and validated alerts.',
     .88, 4.39, 5.43, 21, 'body', 'dim')
text('BUILD → VERIFY → LEARN → OPERATE', .88, 5.73, 5.6, 11, 'mono', 'lilac')
route_plot(7.35, 2.28, 5.0, 3.40, True)
text('LOFOTEN / 4 IMPORTED ROUTES', 7.30, 5.96, 5.1, 10, 'mono', 'lilac')
takeaway('Current position: a working local investigation platform with a clear path to field validation.')

# 02
slide('01 / Executive position', 'The foundation is built. Field validation is the next gate.',
      'Progress spans capture, asset context, investigation and exploratory analytics.',
      'SOURCES / docs/map-app.md · docs/masts.md · docs/event-ml.md · recorded DB snapshot',
      'Counts are the recorded 16 September 2026 snapshot, not a fresh server query. Existing 115,847 legacy events are unclassified threshold records. The 575 connections are inferred. New investigation tables and review writes are local; server class and asset migrations were previously applied.')
card(.55, 2.23, 3.94, 2.30, 'Built', 'Evidence workspace',
     'Map assets, inspect raw windows, replay recordings and record classification reviews.')
card(4.69, 2.23, 3.94, 2.30, 'Demonstrated', 'Discovery workflow',
     'A bounded archive pilot extracts features and ranks candidate windows for expert review.')
card(8.83, 2.23, 3.95, 2.30, 'Next', 'Validated detection',
     'Verify timing and geography; collect independent labels; evaluate by incident and season.')
for x, v, label in [(0.75, '579', 'imported masts'), (3.86, '575', 'inferred connections'),
                    (6.97, '115,847', 'legacy RMS event records'), (10.18, '4', 'target event classes')]:
    stat(x, 4.85, 2.60, v, label)
takeaway('A strong software foundation; classification accuracy and operational benefit remain to be measured.')

# 03
slide('02 / Product intent', 'Turn distributed measurements into actionable evidence',
      'The intended operator outcome: find where to investigate, why it matters and what supports the decision.',
      'SOURCES / AGENTS.md · docs/event-ml.md; operational outcomes are project goals',
      'DAS means distributed acoustic sensing. A channel is a spatial sample along the monitored fibre. The target classes are not the RMS/variance metric identifiers in the database. Vegetation strikes are a candidate subtype of External Interference, not an implemented dedicated classifier. Ice mass, load and displacement cannot be inferred without calibration and ground truth.')
targets = [('01', 'Normal\nOperations', 'Establish the range of normal behaviour, including wind.'),
           ('02', 'Ice\nLoading', 'Test whether signal changes can distinguish ice-related conditions.'),
           ('03', 'Wind\nGalloping', 'Validate sustained motion patterns against field observations.'),
           ('04', 'External\nInterference', 'Review strikes, vegetation contact and other disturbances.')]
for i, (n, t, b) in enumerate(targets):
    card(.55+i*3.11, 2.26, 2.90, 2.95, f'Target class {n}', t, b, 26)
text('MEASURE', .73, 5.68, 1.95, 14, 'mono', 'lilac')
text('→', 2.88, 5.58, .4, 26, 'body', 'purple')
text('LOCATE', 3.53, 5.68, 1.95, 14, 'mono', 'lilac')
text('→', 5.67, 5.58, .4, 26, 'body', 'purple')
text('REVIEW', 6.38, 5.68, 1.90, 14, 'mono', 'lilac')
text('→', 8.48, 5.58, .4, 26, 'body', 'purple')
text('PRIORITISE ACTION', 9.1, 5.68, 3.2, 14, 'mono', 'lilac')
takeaway('Proposed value: earlier investigation and better maintenance prioritisation, measured during the pilot.')

# 04
slide('03 / System architecture', 'One evidence chain connects signals, assets and review',
      'HDF5 stores the signal; PostgreSQL stores context; the API serves bounded analysis to the browser.',
      'SOURCES / docs/map-app.md · src/Mapui/{api,signals,metrics,live}.py',
      'React/TypeScript, MapLibre and Plotly form the frontend. FastAPI serves PostgreSQL metadata and bounded HDF5 slices. Capture accepts completed-file notifications and publishes atomic file-backed snapshots. Replay is implemented and explicitly labeled. Live operation requires independently verified acquisition time and channel mapping; notification support is not evidence of field validation. No complete raw signal array is posted as an in-memory global.')
card(.55, 2.25, 2.60, 2.72, 'Input', 'DAS archive', 'Interrogator\nCompleted HDF5 files\nExplicit channel IDs', 25)
card(3.63, 2.25, 2.68, 2.72, 'Processing', 'Bounded analysis', 'Python / h5py\nFixed-frame metrics\nReplay + file snapshots', 25)
card(6.80, 2.25, 2.68, 2.72, 'Service', 'FastAPI', 'Signal windows\nAsset geometry\nEvents and review API', 25)
card(9.97, 2.25, 2.81, 2.72, 'Operator', 'Cable Observatory', 'React / MapLibre\nPlotly signal review\nClassification history', 25)
for x in [3.23, 6.41, 9.57]:
    text('→', x, 3.27, .40, 26, 'body', 'purple')
rect(3.63, 5.33, 9.15, .65)
text('POSTGRESQL   /   assets · events · weather · recording references · review history',
     3.84, 5.51, 8.70, 12, 'mono', 'lilac')
line(8.14, 4.97, 8.14, 5.33, 'purple', 2)
takeaway('The local investigation workflow is implemented; verified live capture and production rollout are separate gates.')

# 05
slide('04 / Geographic foundation', '579 masts anchor the asset view across four routes',
      'Workbook provenance and channel references are retained for every imported mast.',
      'SOURCES / results/masts/masts.csv · results/masts/spans.csv · docs/masts.md',
      'Diagram generated directly from imported mast coordinates, ordered by route_position; links are illustrative inferred connections, not the actual fibre path. Route names come from the mast workbook inventory, not conflicting event-limit gids. R13 lacks M38. R06 M184–M188 lie 17.5–498.5 m from the endpoint channel; R14 M1 is about 30.1 m away. All other imported masts are within 10 m. Seven water-crossing candidates are unconfirmed and all installation types remain unknown.')
rect(.55, 2.24, 7.12, 3.95)
route_plot(.84, 2.45, 6.48, 3.08)
text('MAST COORDINATES / INFERRED CONNECTIONS / NO BASEMAP', .78, 5.83, 6.65, 9, 'mono', 'dim')
routes = [('R06', 'Fygle–Solbjørn', '188', '187', 'lilac'),
          ('R12', 'Svolvær–Kleppstad', '94', '93', 'purple'),
          ('R13', 'Kvitfossen–Svolvær', '101', '100', 'pink'),
          ('R14', 'Kleppstad–Fygle via Valberg', '196', '195', 'fg')]
text('ROUTE / MASTS / CONNECTIONS', 8.02, 2.28, 4.6, 10, 'mono', 'dim')
for i, (route, name, masts, spans, color) in enumerate(routes):
    yy=2.76+i*.78
    line(8.02, yy+.12, 8.26, yy+.12, color, 3)
    text(route, 8.40, yy-.03, 1.0, 18, 'head', color)
    text(f'{masts} / {spans}', 10.96, yy-.02, 1.7, 17, 'mono')
    text(name, 8.40, yy+.28, 4.15, 13, 'body', 'dim')
takeaway('Assets are mapped; physical connectivity, water crossings and distant endpoint matches still need review.')

# 06
slide('05 / Working application', 'Investigate an event, a place or a recording',
      'An integrated local workspace connects the event queue, network map and signal evidence.',
      'SOURCE / .local/screenshots/overview.png · docs/map-app.md',
      'Unmodified screenshot of the current local application. Its Nord styling and legacy Sensnet label are retained as implementation evidence; the presentation introduces Ubeconn branding separately. The recording shown is an independent archive example, not a verified window for historical events. Map attribution is retained within the screenshot. Basemap: OpenFreeMap / OpenMapTiles / OpenStreetMap. Local review identities are self-reported; authentication is a deployment task.')
image(ROOT / '.local/screenshots/overview.png', .60, 2.14, h=4.61)
card(6.03, 2.23, 6.75, 1.10, 'Locate', 'Search routes, masts and recorded events', '', 23)
card(6.03, 3.55, 6.75, 1.10, 'Inspect', 'Waveform, heatmap, spectrum and trends', '', 23)
card(6.03, 4.87, 6.75, 1.10, 'Review', 'Store a class, evidence note and history', '', 23)
text('ACTUAL LOCAL BUILD / ORIGINAL UI COLOURS', 6.24, 6.32, 6.2, 10, 'mono', 'lilac')

# 07
slide('06 / Signal integrity', 'Signal amplitude and event meaning stay distinct',
      'Map layers are reproducible window measurements. Classification requires independent evidence.',
      'SOURCES / docs/map-app.md · docs/channel-mapping.md · src/Mapui/metrics.py',
      'The local app applies header dataScale and preserves explicit header channel IDs. Fixed ten-second frames make map metrics independent of display zoom. AC RMS and absolute peak remove the window mean; band RMS integrates Welch power, with exploratory default 1–10 Hz; variance uses squared units. Display reduction retains actual extrema. Current example recordings resolve 17,010 exact positions and 8,224 bounded coordinate interpolations. Two channels outside the stored mapping remain unlocated. No signal values are interpolated or channel IDs shifted.')
for i, (t, b) in enumerate([('AC RMS', 'Vibration level after removing the mean'),
                           ('Band RMS', 'Energy within a selected frequency band'),
                           ('Absolute peak', 'Largest mean-removed signal excursion'),
                           ('Variance', 'Signal variability in squared units')]):
    card(.55+i*3.11, 2.24, 2.9, 2.0, 'Map metric', t, b, 24)
text('TRACEABILITY', .72, 4.67, 2.30, 11, 'mono', 'lilac')
text('File + sample window + channel ID + scale + processing version', 3.17, 4.57, 9.32, 24, 'head')
for x, v, label in [(0.74, '17,010', 'exact coordinate matches'), (4.84, '8,224', 'bounded coordinate interpolations'),
                    (9.24, '2', 'channels remain unlocated')]:
    text(v, x, 5.17, 3.6, 30, 'head', 'lilac')
    text(label, x, 5.78, 3.6, 12, 'body', 'dim')
takeaway('Never interpret map colour as event class or severity; no missing fibre ranges are reconstructed.')

# 08
slide('07 / Evidence readiness', 'The largest archive is still an unlabeled candidate set',
      'Time, location and labels must be verified together before historical events become training evidence.',
      'SOURCES / AGENTS.md · docs/map-app.md · docs/event-ml.md · .local/verification.json',
      '115,847 records span 22 April–2 June 2025 and are all RMS metric events with NULL class in the inspected snapshot. Their comments do not match current threshold names, so the legacy location join cannot resolve them. Current channel metadata is dated 16 April 2026. Example archive headers suggest 16 June 2026 but original acquisition-time interpretation remains unverified. A 2026 recording cannot be attached to a 2025 event merely for visual completeness. NULL means unclassified, never normal.')
rows = [
    ('Legacy alarms', 'APR–JUN 2025', '115,847 RMS records', 'Class and reliable event location unresolved'),
    ('Channel metadata', 'APR 2026', '25,235 positions', 'Historical mapping equivalence unverified'),
    ('Example recordings', 'JUN 2026*', '3 full local files', 'Independent of the historical alarm archive'),
    ('Expert ground truth', 'NEXT PHASE', 'Incident / inspection evidence', 'Required to validate the four target classes'),
]
for i, (name, date, amount, caveat) in enumerate(rows):
    yy=2.25+i*.90
    if i%2==0: rect(.55, yy, 12.23, .84)
    text(name, .76, yy+.13, 2.60, 20, 'head')
    text(date, 3.47, yy+.17, 1.86, 11, 'mono', 'lilac')
    text(amount, 5.55, yy+.14, 2.67, 16, 'medium')
    text(caveat, 8.55, yy+.12, 3.91, 15, 'body', 'dim')
text('* Header date is plausible but not independently verified.', .75, 6.04, 10, 11, 'body', 'dim')
takeaway('Alarm volume is useful for review prioritisation; it does not establish labeled incidents or model accuracy.')

# 09
slide('08 / Weather context', 'Weather adds context once time and location are trustworthy',
      'YR collection is in place at two configured locations; humidity is a planned extension.',
      'SOURCES / src/weather/weather_capture.py · docs/map-app.md · docs/event-discovery.md',
      'Stored variables are air temperature, wind speed and wind direction at Kråkern Bru and Leknes. These are collected YR forecast values, not necessarily direct measurements. The current collector uses database-default collection timestamps rather than verified forecast-validity times. The local UI only uses past rows relative to the stored event time and displays missing/stale context; this does not independently verify the original acquisition alignment. Add humidity with explicit lookup semantics, units, collector support and missing-value handling. Do not assume numeric type 4 means humidity.')
card(.55, 2.25, 5.95, 2.65, 'Collected today', 'Temperature · wind speed · direction',
     'Kråkern Bru and Leknes\nYR forecast context\nMissing and stale values remain visible', 25)
card(6.74, 2.25, 6.04, 2.65, 'Proposed extension', 'Humidity + source-validity timestamps',
     'Explicit variable and unit definitions\nCollector and storage support\nDocumented station-to-route relevance', 25)
text('ALIGNMENT RULE', .77, 5.06, 3.0, 11, 'mono', 'lilac')
text('Verified acquisition time → appropriate location → latest valid prior weather',
     .77, 5.49, 11.6, 24, 'head')
takeaway('Weather supports interpretation; humidity or freezing conditions alone do not prove ice loading.')

# 10
slide('09 / Discovery pilot', 'Archive discovery has been exercised on real recordings',
      'A reproducible pilot produces a compact review queue without loading the full archive into memory.',
      'SOURCES / results/das-discovery-pilot-20260916/{summary,pilot-validation}.json',
      'Pilot inputs: 90 approximately ten-second files, 25,236 channels, roughly 15 minutes of archive data. Fifteen 60-second temporal windows × 16 spatial blocks × 32 channels per block produce 240 feature rows. Six exploratory clusters have populations 45, 12, 23, 131, 7, 22 for IDs 0–5. Initial run 37.849 seconds, resume 0.288 seconds; max RSS 164300 KiB = approximately 160.4 MiB. Resumed feature output had identical SHA-256. These are one recorded run, not general performance guarantees. No weather join was enabled; timestamps and locations were unverified. Scores are review priorities, not class probabilities.')
for x, v, label in [(.76, '90', 'archive files'), (3.85, '15 min', 'approximate coverage'),
                    (6.93, '240', 'window / block feature rows'), (10.11, '6', 'exploratory clusters')]:
    stat(x, 2.27, 2.70, v, label)
text('CLUSTER POPULATION / FEATURE ROWS', .76, 3.78, 6.1, 10, 'mono', 'lilac')
counts = [45,12,23,131,7,22]
for i, count in enumerate(counts):
    yy=4.18+i*.29
    text(str(i), .78, yy-.025, .33, 10, 'mono', 'dim')
    rect(1.20, yy, count/131*4.42, .17, 'purple' if i!=3 else 'lilac')
    text(str(count), 1.31+count/131*4.42, yy-.036, .6, 10, 'mono')
card(7.06, 3.80, 5.72, 2.31, 'Recorded execution', '37.8 s initial / 0.29 s resume',
     '≈160 MiB peak resident memory\nIdentical feature output after resume\nOne pilot run; no accuracy estimate', 24)
takeaway('This proves the discovery workflow on a small archive; seasonal coverage and verified labels are still needed.')

# 11
slide('10 / Detection development', 'The ML baseline is built; field accuracy is still unproven',
      'The implementation supports four classes, grouped validation and quantified prediction windows.',
      'SOURCE / docs/event-ml.md · src/Capture/event_ml.py',
      'No trained supervised model or measured field accuracy is supplied. The offline baseline uses stored input values, while discovery and map analysis apply dataScale; preprocessing must be reconciled before model reuse. At least three independent groups per class are a code-level validation minimum, not sufficient field evidence. Group related windows by session or incident to avoid leakage. Hold out sites, seasons and incidents. Current output probabilities are uncalibrated, and the default 0.7 review threshold is not an approved operating point. The single-label model cannot describe simultaneous ice and galloping.')
stages=[('01', 'Discover', 'Rank unusual and representative windows.'),
        ('02', 'Verify', 'Attach inspection, camera or incident evidence.'),
        ('03', 'Train', 'Use one agreed signal representation.'),
        ('04', 'Evaluate', 'Hold out independent incidents and seasons.')]
for i, (n,t,b) in enumerate(stages):
    card(.55+i*3.11, 2.27, 2.9, 2.40, n, t, b, 27)
rect(.55, 4.99, 12.23, 1.08)
text('RELEASE EVIDENCE', .78, 5.20, 2.68, 11, 'mono', 'lilac')
text('Per-class precision / recall · false-alert burden · confidence calibration\nSite and season generalisation · detection latency · traceable predictions',
     3.61, 5.17, 8.84, 16, 'body')
takeaway('Model deployment follows field evidence; threshold alarms and anomaly clusters cannot substitute for labels.')

# 12
slide('11 / Readiness assessment', 'Software checks do not establish field performance',
      'Recorded local checks support the investigation workflow; operational acceptance remains ahead.',
      'SOURCES / docs/map-app.md · docs/event-ml.md · docs/masts.md · recorded pilot results',
      'The documentation records 58 Python tests passed, a successful production frontend build, desktop/mobile browser checks with no page errors, save/reload/history checks, and three recordings published as archive replay on 16 September 2026. These checks were not rerun for this presentation and do not establish a live field deployment. Asset migrations 001–003 were previously applied to the server. Investigation migration 004 is in the local clone only; production deployment is not included.')
headers=[('Capability', .76, 3.5), ('Evidence available', 4.48, 4.4), ('Remaining gate', 9.05, 3.3)]
for s,x,w in headers: text(s.upper(), x, 2.23, w, 10, 'mono', 'lilac')
items=[('Asset foundation', '579 masts / 575 inferred links', 'Confirm physical inventory'),
       ('Investigation app', 'Browser workflow + review persistence', 'Authenticated multi-user deployment'),
       ('Signal and replay', 'Bounded reads / 3 files replayed', 'Verified live timing and mapping'),
       ('Discovery', '240 feature rows / repeatable resume', 'Broader seasonal data'),
       ('Four-class baseline', 'Training + grouped evaluation code', 'Labels and held-out field results')]
for i,(a,b,c) in enumerate(items):
    yy=2.72+i*.61
    if i%2==0: rect(.55, yy-.04, 12.23, .59)
    text(a, .76, yy+.04, 3.46, 17, 'medium')
    text(b, 4.48, yy+.04, 4.27, 16, 'body', 'dim')
    text(c, 9.05, yy+.04, 3.43, 15, 'body', 'dim')
text('RECORDED LOCAL CHECKS: 58 PYTHON TESTS + FRONTEND BUILD + BROWSER REVIEW CHECKS',
     .77, 6.02, 11.90, 10, 'mono', 'lilac')
takeaway('The next milestone is a controlled field pilot with measurable acceptance criteria.')

# 13
slide('12 / Proposed roadmap', 'Advance through evidence gates, then expand deployment',
      'Sequence proposed for planning; durations, owners and budget require project agreement.',
      'PROPOSAL / grounded in the documented project limitations and deployment prerequisites',
      'This is a proposed execution sequence, not an approved schedule. Workstreams can overlap when dependencies are satisfied. Ice-related validation depends on observing representative conditions, so calendar promises would be premature. Agree accountable owners and gate criteria before execution. Phase 1 includes resolving legacy mapping conflicts and an explicit preprocessing contract; phase 2 adds humidity and source-validity timestamps; phase 3 validates on independent incidents; phase 4 begins in shadow mode with human review.')
stages=[('GATE 1', 'Trust the data', 'Verify acquisition time, channel mapping and preprocessing.',
         'Exit: traceable reference windows'),
        ('GATE 2', 'Build ground truth', 'Collect field labels, normal baselines and richer weather context.',
         'Exit: reviewed, grouped dataset'),
        ('GATE 3', 'Prove detection', 'Evaluate held-out incidents; calibrate thresholds and confidence.',
         'Exit: agreed performance achieved'),
        ('GATE 4', 'Pilot operations', 'Run shadow alerts, measure workload and harden deployment.',
         'Exit: operator acceptance')]
for i,(n,t,b,g) in enumerate(stages):
    xx=.55+i*3.11
    card(xx,2.30,2.9,3.54,n,t,b,26)
    line(xx+.2,4.83,xx+2.7,4.83)
    text(g,xx+.2,5.04,2.49,15,'medium','lilac')
takeaway('Protect the evidence chain first; expand sites, seasons and automation only after acceptance gates pass.')

# 14
slide('13 / Proposed acceptance scorecard', 'Measure detection quality and operational usefulness',
      'Agree numerical targets before the pilot; the project does not yet have validated operating thresholds.',
      'PROPOSAL / validation plan; no performance results or target commitments implied',
      'Report performance by class and condition with sample counts and uncertainty. Evaluate at incident level, not only window level. False alerts should be reported per monitored route-day or another agreed coverage unit. Acquisition-to-alert latency must distinguish acquisition time from publication and processing time. Location error requires verified physical events. Measure missing data and unavailable coverage. Pair technical results with review burden, useful findings and independently verified outcomes; avoided costs remain unproven.')
rows=[('Data integrity', 'Verified timestamps, channel coverage, missing-data rate', 'Acquisition + data owner'),
      ('Event quality', 'Per-class precision / recall; false alerts per route-day', 'Analytics + domain reviewer'),
      ('Location and latency', 'Error against verified location; acquisition-to-alert delay', 'Field + platform owner'),
      ('Operational value', 'Review time, actionable findings, missed confirmed incidents', 'Operations owner'),
      ('Robustness', 'Performance across sites, seasons and configuration changes', 'Project acceptance team')]
for t,x,w in [('MEASURE',.76,2.65),('PILOT EVIDENCE',3.58,5.77),('PROPOSED OWNER',9.62,2.95)]:
    text(t,x,2.27,w,10,'mono','lilac')
for i,(a,b,c) in enumerate(rows):
    yy=2.77+i*.66
    if i%2==0:rect(.55,yy-.04,12.23,.65)
    text(a,.76,yy+.06,2.60,18,'head')
    text(b,3.58,yy+.04,5.72,16,'body','dim')
    text(c,9.62,yy+.04,2.9,15,'body','dim')
takeaway('Pilot success means trusted evidence and a manageable alert burden under representative field conditions.')

# 15
slide('14 / Production path', 'Prepare the platform for controlled operations',
      'A production pilot adds identity, reliability and incident handling around the existing evidence workflow.',
      'SOURCES / docs/map-app.md deployment boundary · docs/event-ml.md next steps; extensions proposed',
      'Current local deployment uses loopback, a local database copy, same-origin write checks and self-reported reviewer names. Production requires authenticated identities, suitable authorization, a restricted database role, reviewed migration 004 and deliberate service enablement. Supervisor templates default autostart=false. Proposed reliability work includes retention planning, backup/restore tests, health monitoring and integration-specific alert delivery. Incident aggregation should preserve source file/sample/channel evidence, distinguish a candidate from a sustained incident and support review/closure.')
card(.55,2.26,3.94,3.25,'Identity + governance','Accountable decisions',
     'Authenticated reviewers\nRestricted database access\nReviewed schema migration\nVersioned models and thresholds',26)
card(4.69,2.26,3.94,3.25,'Reliability','Observable operation',
     'Capture health and stale-data status\nStorage and retention planning\nBackup / restore verification\nLive / replay separation',26)
card(8.83,2.26,3.95,3.25,'Incident workflow','Useful alerts',
     'Merge related candidate windows\nPreserve evidence references\nHuman review and closure\nAgreed escalation integrations',26)
text('SHADOW MODE FIRST', .78,5.91,3.21,12,'mono','lilac')
text('Observe decisions and workload before automated operational escalation.',4.03,5.83,8.43,19,'body')
takeaway('Deploy a support tool with traceable evidence; increase automation as field confidence grows.')

# 16
slide('15 / Executive decisions', 'Fund the next evidence milestone',
      'The proposed next increment is a controlled validation pilot with named owners and a reviewed dataset.',
      'PROPOSAL / scope, resources, target values and schedule to be agreed by Ubeconn AS',
      'No investment amount, commercial claim, customer commitment or staffing assignment is implied. Proposed decisions: select representative pilot routes and field access; assign data, domain-review and operations accountability; approve acceptance measures and production prerequisites. The concrete deliverable is a verified evidence pack plus held-out results and an operational acceptance decision. Data capture should include quiet normal operation, windy normal operation and independently verified target incidents. Seasonal or rare conditions may require a longer observation period.')
card(.55,2.27,3.94,2.65,'Decision 1','Choose the pilot scope',
     'Representative routes, field access and independently observable conditions.',28)
card(4.69,2.27,3.94,2.65,'Decision 2','Assign the evidence owners',
     'Acquisition and mapping, expert labeling, analytics and operator review.',28)
card(8.83,2.27,3.95,2.65,'Decision 3','Agree the acceptance gate',
     'Detection quality, alert workload, traceability and deployment readiness.',28)
text('NEXT DELIVERABLE',.77,5.35,3.05,11,'mono','lilac')
text('Verified dataset + held-out evaluation + pilot go / no-go',3.96,5.22,8.44,27,'head')
takeaway('Ubeconn’s next step: convert a working technical foundation into defensible operational evidence.')

# 17
slide('Appendix A / Evidence register', 'Sources and scope behind the presentation',
      'Project-local evidence only. No external market, cost-saving or detection-accuracy claims are made.',
      'COMPILED / 16 SEP 2026 · repository working tree and recorded project evidence',
      'Source precedence: current implementation documentation and artifacts supersede older architecture descriptions in AGENTS.md. Specifically the capture pipeline now uses file-backed snapshots and the main application is FastAPI + React rather than the retired Dash contract. Snapshot counts are historical observations, not refreshed server queries. No database, application or production configuration was changed to create this presentation. Roadmap and acceptance slides are proposals.')
refs=[('Platform and deployment','docs/map-app.md; src/Mapui/{api,signals,metrics,live}.py'),
      ('Asset inventory and mapping','docs/masts.md; docs/channel-mapping.md; results/masts/*.csv'),
      ('Database snapshot','AGENTS.md; .local/verification.json (16 September 2026)'),
      ('Discovery pilot','docs/event-discovery.md; results/das-discovery-pilot-20260916/'),
      ('Supervised baseline','docs/event-ml.md; src/Capture/event_ml.py'),
      ('Application image','Original .local/screenshots/overview.png; map attribution retained')]
for i,(a,b) in enumerate(refs):
    yy=2.27+i*.56
    text(a,.77,yy,3.40,17,'head')
    text(b,4.34,yy+.04,8.16,12,'mono','dim')
text('READING THE STATUS',.77,5.91,3.26,10,'mono','lilac')
text('Implemented / field validated\nInferred / confirmed     Proposed / committed',4.34,5.84,8.12,16,'medium')
takeaway('All numerical project claims refer to the documented snapshot or recorded pilot run.')

# 18
slide('Appendix B / Ubeconn visual identity', 'A simple identity, grounded in ube purple',
      'House-report hierarchy and IBM Plex typography, adapted to Ubeconn’s colour profile.',
      'BRAND ASSETS / white vector wordmark + transparent PNG + IBM Plex fonts (OFL)',
      'The logo is intentionally a plain lowercase white wordmark. The colours are a proposed Ubeconn palette inspired by ube flesh, not a standardized botanical colour reference. Background: deep plum #1D1429. Panel #2B1E3A. Primary ube #9D6BD0. Accent lavender #CFADF0. Text #F6F0FC. Plain logo #FFFFFF. IBM Plex Sans is body text, IBM Plex Sans Condensed bold is for headings, IBM Plex Mono is for labels. Fonts are supplied under the included SIL Open Font License. PowerPoint text remains editable; the PDF is a font-stable companion.')
rect(.55,2.25,5.02,3.75)
text('ubeconn',.99,3.12,4.15,68,'medium','white')
text('PLAIN WHITE WORDMARK',1.02,4.75,4.17,11,'mono','lilac')
swatches=[('Deep plum','1D1429'),('Ube','9D6BD0'),('Lavender','CFADF0'),('Soft white','F6F0FC')]
for i,(n,c) in enumerate(swatches):
    xx=5.94+i*1.72
    rect(xx,2.25,1.55,1.37,c)
    text(n,xx,3.85,1.57,15,'medium')
    text('#'+c,xx,4.23,1.57,10,'mono','dim')
text('IBM Plex Sans Condensed',5.94,4.92,6.68,27,'head')
text('IBM Plex Sans / Body     IBM Plex Mono / Labels',5.96,5.57,6.70,14,'body','dim')
takeaway('Dark backgrounds, generous spacing, restrained colour and clear evidence labels.')


def rgb(s):
    return tuple(int(s[i:i+2],16) for i in (0,2,4))

def write_pptx():
    prs=Presentation()
    prs.slide_width=Inches(W)
    prs.slide_height=Inches(H)
    prs.core_properties.title='Ubeconn AS | Cable intelligence — executive project review'
    prs.core_properties.subject='Project progress, technical evidence and proposed validation roadmap'
    prs.core_properties.author='Ubeconn AS'
    prs.core_properties.keywords='Ubeconn, DAS, cable monitoring, executive review, September 2026'
    prs.core_properties.comments='Editable slide content. Source details and caveats are in speaker notes.'
    for d in SLIDES:
        s=prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb=RGBColor.from_string(C['bg'])
        for o in d['objects']:
            kind=o['kind']
            if kind in ('rect','circle'):
                shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE if kind=='rect' else MSO_SHAPE.OVAL,
                                        Inches(o['x']),Inches(o['y']),Inches(o['w']),Inches(o['h']))
                shape.fill.solid()
                shape.fill.fore_color.rgb=RGBColor.from_string(o['color'])
                shape.line.fill.background()
            elif kind=='line':
                shape=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(o['x']),Inches(o['y']),
                                            Inches(o['x2']),Inches(o['y2']))
                shape.line.color.rgb=RGBColor.from_string(o['color'])
                shape.line.width=Pt(o['width'])
            elif kind=='image':
                s.shapes.add_picture(o['path'],Inches(o['x']),Inches(o['y']),Inches(o['w']),Inches(o['h']))
            elif kind=='text':
                shape=s.shapes.add_textbox(Inches(o['x']),Inches(o['y']),Inches(o['w']),Inches(o['h']))
                tf=shape.text_frame
                tf.clear()
                tf.word_wrap=False
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                for j,txt in enumerate(o['text'].split('\n')):
                    p=tf.paragraphs[0] if j==0 else tf.add_paragraph()
                    p.text=txt
                    p.space_before=Pt(0)
                    p.space_after=Pt(0)
                    p.line_spacing=Pt(o['size']*1.22)
                    p.font.name=FONT_SPECS[o['role']][2]
                    p.font.size=Pt(o['size'])
                    p.font.bold=o['role'] in ('head','medium')
                    p.font.color.rgb=RGBColor.from_string(o['color'])
        s.notes_slide.notes_text_frame.text=d['notes']+'\n\nSources: '+d['source']
    # Theme: Ubeconn colours and IBM Plex, so new editable objects inherit the brand.
    theme=prs.slide_masters[0].part.theme_part if hasattr(prs.slide_masters[0].part,'theme_part') else None
    from lxml import etree
    for part in prs.part.package.iter_parts():
        if str(part.partname).startswith('/ppt/theme/theme'):
            tree=etree.fromstring(part.blob)
            ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
            clr=tree.find('.//a:clrScheme',ns)
            clr.set('name','Ubeconn')
            scheme=['1D1429','F6F0FC','2B1E3A','FFFFFF','9D6BD0','CFADF0','DFAFD2','BAA9CB','735096','493458','CFADF0','DFAFD2']
            for node,color in zip(clr,scheme):
                for child in list(node):node.remove(child)
                ch=etree.SubElement(node,'{'+ns['a']+'}srgbClr');ch.set('val',color)
            tree.find('.//a:majorFont/a:latin',ns).set('typeface','IBM Plex Sans Condensed')
            tree.find('.//a:minorFont/a:latin',ns).set('typeface','IBM Plex Sans')
            part._blob=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    prs.save(OUT/'Ubeconn_Executive_Review_2026-09.pptx')


def write_previews_and_pdf():
    scale=144 # 1920 × 1080 preview
    pdf=canvas.Canvas(str(OUT/'Ubeconn_Executive_Review_2026-09.pdf'),pagesize=(W*72,H*72))
    pdf.setTitle('Ubeconn AS | Executive project review | September 2026')
    pdf.setAuthor('Ubeconn AS')
    for index,d in enumerate(SLIDES,1):
        im=Image.new('RGB',(round(W*scale),round(H*scale)),rgb(C['bg']))
        draw=ImageDraw.Draw(im)
        pdf.setFillColorRGB(*(v/255 for v in rgb(C['bg'])))
        pdf.rect(0,0,W*72,H*72,fill=1,stroke=0)
        for o in d['objects']:
            k=o['kind'];x=o['x'];y=o['y']
            if k in ('rect','circle'):
                box=(round(x*scale),round(y*scale),round((x+o['w'])*scale),round((y+o['h'])*scale))
                pdf.setFillColorRGB(*(v/255 for v in rgb(o['color'])))
                if k=='rect':
                    draw.rectangle(box,fill=rgb(o['color']))
                    pdf.rect(x*72,(H-y-o['h'])*72,o['w']*72,o['h']*72,fill=1,stroke=0)
                else:
                    draw.ellipse(box,fill=rgb(o['color']))
                    pdf.ellipse(x*72,(H-y-o['h'])*72,(x+o['w'])*72,(H-y)*72,fill=1,stroke=0)
            elif k=='line':
                draw.line((round(x*scale),round(y*scale),round(o['x2']*scale),round(o['y2']*scale)),
                          fill=rgb(o['color']),width=max(1,round(o['width']*scale/72)))
                pdf.setStrokeColorRGB(*(v/255 for v in rgb(o['color'])))
                pdf.setLineWidth(o['width'])
                pdf.line(x*72,(H-y)*72,o['x2']*72,(H-o['y2'])*72)
            elif k=='image':
                pic=Image.open(o['path']).convert('RGB').resize((round(o['w']*scale),round(o['h']*scale)),Image.Resampling.LANCZOS)
                im.paste(pic,(round(x*scale),round(y*scale)))
                pdf.drawImage(o['path'],x*72,(H-y-o['h'])*72,o['w']*72,o['h']*72)
            elif k=='text':
                f=ImageFont.truetype(str(FONT_PATHS[o['role']]),round(o['size']*scale/72))
                pdf.setFont(o['role'],o['size'])
                pdf.setFillColorRGB(*(v/255 for v in rgb(o['color'])))
                for j,t in enumerate(o['text'].split('\n')):
                    baseline=y*scale+(o['size']*.97+j*o['size']*1.22)*scale/72
                    draw.text((x*scale,baseline),t,font=f,anchor='ls',fill=rgb(o['color']))
                    pdf.drawString(x*72,(H-y)*72-o['size']*.97-j*o['size']*1.22,t)
        im.save(PREVIEWS/f'slide-{index:02d}.png')
        pdf.showPage()
    pdf.save()
    thumbw=480;thumbh=270
    sheet=Image.new('RGB',(thumbw*3,thumbh*6),(15,10,21))
    for i in range(len(SLIDES)):
        im=Image.open(PREVIEWS/f'slide-{i+1:02d}.png').resize((thumbw,thumbh),Image.Resampling.LANCZOS)
        sheet.paste(im,((i%3)*thumbw,(i//3)*thumbh))
    sheet.save(OUT/'slide-overview.png')


def validate():
    issues=[]
    for i,d in enumerate(SLIDES,1):
        for o in d['objects']:
            if o['kind']=='line':continue
            if o['x'] < -.01 or o['y'] < -.01 or o['x']+o['w'] > W+.01 or o['y']+o['h'] > H+.01:
                issues.append((i,'outside slide',o.get('text','')))
            if o['kind']=='text':
                for txt in o['text'].split('\n'):
                    font=ImageFont.truetype(str(FONT_PATHS[o['role']]),round(o['size']*4))
                    if font.getlength(txt) > o['w']*72*4:
                        issues.append((i,'text overflow',txt))
                # Header and source text excepted: all body text must stop above the footer.
                if 2.05<o['y']<6.45 and o['y']+o['h']>6.95:
                    issues.append((i,'body reaches footer',o['text']))
    if issues:raise RuntimeError(issues)
    p=Presentation(OUT/'Ubeconn_Executive_Review_2026-09.pptx')
    assert len(p.slides)==18
    assert all(s.has_notes_slide for s in p.slides)
    (OUT/'validation.json').write_text(json.dumps({'slides':len(p.slides),'bounds_and_text_width':'passed',
        'speaker_notes':'all slides','preview_method':'shared layout rendered with bundled fonts; not an Office-engine render',
        'source_masts':len(MASTS),'pilot_feature_rows':SUMMARY['feature_rows']},indent=2))

write_pptx()
validate()
write_previews_and_pdf()
(OUT/'speaker-notes.md').write_text('\n\n'.join(f'## {i:02d} — {d["title"]}\n\n{d["notes"]}\n\n{d["source"]}' for i,d in enumerate(SLIDES,1)))
(OUT/'brand-palette.json').write_text(json.dumps(C,indent=2))
print(f'Wrote {len(SLIDES)} slides to {OUT}')
